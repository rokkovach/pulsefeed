"""Build the editable six-page report from review.md and reviewed PNG assets.

The PDF is exported from this DOCX so the two deliverables contain the same text.
Run the document skill's renderer after building; generated page images are QA
artifacts and do not belong in the public handoff.
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(__file__).with_name('review.md')
OUTPUT = ROOT / 'output/docx/pulsefeed-review.docx'


def rich(paragraph, text):
    """Keep emphasis editable instead of baking narrative into an image."""
    for index, part in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        run = paragraph.add_run(part)
        run.bold = bool(index % 2)


def table(document, rows):
    result = document.add_table(rows=0, cols=len(rows[0]))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    widths = [1.15, 5.7] if len(rows[0]) == 2 else [1.0, 2.85, 3.0]
    for column, width in zip(result.columns, widths):
        column.width = Inches(width)
    for index, values in enumerate(rows):
        row = result.add_row()
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                margin = OxmlElement('w:' + side); margin.set(qn('w:w'), '100'); margin.set(qn('w:type'), 'dxa'); margins.append(margin)
            props.append(margins)
            borders = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                border = OxmlElement('w:' + side); border.set(qn('w:val'), 'single'); border.set(qn('w:sz'), '5'); border.set(qn('w:color'), 'D9D9D9'); borders.append(border)
            props.append(borders)
            fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'E3EADD' if index == 0 else 'FFFFFF' if index % 2 else 'F6F7F4'); props.append(fill)
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.06
            run = p.add_run(value); run.font.size = Pt(9.3); run.bold = index == 0
        if index == 0:
            repeat = OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
        # Each row must stay readable; no fixed heights or split specification rows.
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def build():
    document = Document()
    # The bundled Word template may carry a blue Title paragraph rule. Remove
    # inherited borders rather than relying on a viewer's theme defaults.
    for style in document.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    section = document.sections[0]
    section.page_width = Inches(8.27); section.page_height = Inches(11.69)
    section.top_margin = Inches(.62); section.bottom_margin = Inches(.62)
    section.left_margin = Inches(.7); section.right_margin = Inches(.7)
    normal = document.styles['Normal']
    normal.font.name = 'Arial'; normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string('242B27')
    normal.paragraph_format.line_spacing = 1.13
    normal.paragraph_format.space_after = Pt(8)
    for name, size in [('Title', 26), ('Heading 1', 22), ('Heading 2', 13)]:
        style = document.styles[name]
        style.font.name = 'Arial'; style.font.size = Pt(size); style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10 if name == 'Heading 2' else 0)
        style.paragraph_format.space_after = Pt(8)
    caption = document.styles['Caption']
    caption.font.name = 'Arial'; caption.font.size = Pt(8.2); caption.font.italic = False; caption.font.bold = False
    caption.font.color.rgb = RGBColor.from_string('57605A')
    caption.paragraph_format.space_after = Pt(8)
    # Quiet page numbering helps a reviewer refer to a worker or screenshot.
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run('PulseFeed  |  '); run.font.name = 'Arial'; run.font.size = Pt(8)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    document.core_properties.title = 'PulseFeed Runpod case study'
    document.core_properties.author = 'Rok Kovac'
    document.core_properties.subject = 'Architecture application model workers validation and next steps'
    document.core_properties.keywords = 'PulseFeed RSS Runpod Railway'
    first_title = True
    for page in SOURCE.read_text().split('<!-- page -->'):
        if not first_title: document.add_page_break()
        lines = page.strip().splitlines(); index = 0
        while index < len(lines):
            line = lines[index].strip()
            if not line: index += 1; continue
            if line.startswith('|'):
                rows = []
                while index < len(lines) and lines[index].strip().startswith('|'):
                    values = [value.strip() for value in lines[index].strip().strip('|').split('|')]
                    if not all(re.fullmatch(r'[: -]+', value) for value in values): rows.append(values)
                    index += 1
                table(document, rows); continue
            if line.startswith('# '):
                document.add_paragraph(line[2:], style='Title' if first_title else 'Heading 1')
                first_title = False
            elif line.startswith('## '): document.add_paragraph(line[3:], style='Heading 2')
            elif line.startswith('!['):
                match = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
                asset = SOURCE.parent / match.group(2)
                p = document.add_paragraph(); p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_with_next = True
                shape = p.add_run().add_picture(str(asset), width=Inches(6.85))
                shape._inline.docPr.set('descr', match.group(1))
                document.add_paragraph(match.group(1), style='Caption')
            elif line.startswith('- '): rich(document.add_paragraph(style='List Bullet'), line[2:])
            else: rich(document.add_paragraph(), line)
            index += 1
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print('Built', OUTPUT.relative_to(ROOT))


if __name__ == '__main__': build()
